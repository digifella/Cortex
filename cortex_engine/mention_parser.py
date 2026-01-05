"""
Mention Parser - Parse @mention Syntax
Version: 1.0.0
Date: 2026-01-05

Purpose: Parse @mention syntax from tender documents and extract:
- Mention type (simple field, structured access, content generation)
- Field path (dot notation)
- Parameters (key=value pairs)

Syntax Examples:
- @companyname                                → Simple field
- @insurance.public_liability.coverage        → Structured access
- @cv[paul_smith, format=brief]              → Content generation with parameters
- @generate[type=executive_summary, topic=digital health, length=500]  → LLM generation
"""

import re
from typing import Optional, Dict, List, Tuple
from enum import Enum
from dataclasses import dataclass

from .utils import get_logger

logger = get_logger(__name__)


class MentionType(str, Enum):
    """Type of mention."""
    SIMPLE = "simple"                  # @companyname
    STRUCTURED = "structured"          # @insurance.public_liability.coverage
    CONTENT_GEN = "content_gen"        # @cv[paul_smith, format=brief]
    CREATIVE_GEN = "creative_gen"      # @generate[type=executive_summary, ...]
    NARRATIVE = "narrative"            # @narrative[company_overview]


@dataclass
class ParsedMention:
    """Parsed mention with all components."""
    raw_text: str                      # Original mention text
    mention_type: MentionType          # Type of mention
    field_path: str                    # Field path (e.g., "insurance.public_liability.coverage")
    parameters: Dict[str, str]         # Parameters from [...] if present
    is_valid: bool                     # Whether parsing succeeded
    error_message: Optional[str] = None

    @property
    def root_field(self) -> str:
        """Get root field name (first part of path)."""
        return self.field_path.split('.')[0] if self.field_path else ""

    @property
    def sub_path(self) -> Optional[str]:
        """Get sub-path after root field."""
        parts = self.field_path.split('.')
        return '.'.join(parts[1:]) if len(parts) > 1 else None

    def __str__(self) -> str:
        """String representation."""
        return f"<Mention: {self.raw_text} | Type: {self.mention_type} | Path: {self.field_path}>"


class MentionParser:
    """Parser for @mention syntax."""

    # Regex patterns
    MENTION_PATTERN = re.compile(
        r'@([a-zA-Z_][a-zA-Z0-9_.]*)(?:\[([^\]]+)\])?',
        re.IGNORECASE
    )

    # Simple field mappings (shorthand → full path)
    SIMPLE_FIELD_MAPPINGS = {
        # Company fields
        'companyname': 'company.legal_name',
        'legalname': 'company.legal_name',
        'abn': 'company.abn',
        'acn': 'company.acn',

        # Contact fields
        'email': 'contact.email',
        'phone': 'contact.phone',
        'website': 'contact.website',
        'registered_office': 'contact.registered_office',
        'registeredoffice': 'contact.registered_office',

        # Special content types
        'cv': 'team.{person_id}',              # Requires person_id parameter
        'project_summary': 'projects.{project_id}',  # Requires project_id parameter
        'reference': 'references.{reference_id}',    # Requires reference_id parameter
        'insurance_summary': 'insurance',      # All insurance
        'team_qualifications': 'team',         # All team qualifications
        'references': 'references',            # All references
        'narrative': 'narrative.{section}',    # Requires section parameter
    }

    # Creative generation types
    CREATIVE_TYPES = {
        'generate',
        'executive_summary',
        'approach',
        'methodology',
        'risk_mitigation',
        'innovation',
        'team_introduction',
        'response'
    }

    def __init__(self):
        """Initialize parser."""
        logger.info("MentionParser initialized")

    def parse(self, mention_text: str) -> ParsedMention:
        """
        Parse a mention string.

        Args:
            mention_text: Raw mention text (e.g., "@companyname" or "@cv[paul_smith]")

        Returns:
            ParsedMention object

        Examples:
            >>> parser = MentionParser()
            >>> parser.parse("@companyname")
            <Mention: @companyname | Type: simple | Path: company.legal_name>

            >>> parser.parse("@insurance.public_liability.coverage")
            <Mention: @insurance... | Type: structured | Path: insurance.public_liability.coverage>

            >>> parser.parse("@cv[paul_smith, format=brief]")
            <Mention: @cv[...] | Type: content_gen | Path: team.paul_smith>
        """
        # Remove whitespace
        mention_text = mention_text.strip()

        # Validate starts with @
        if not mention_text.startswith('@'):
            return ParsedMention(
                raw_text=mention_text,
                mention_type=MentionType.SIMPLE,
                field_path="",
                parameters={},
                is_valid=False,
                error_message="Mention must start with @"
            )

        # Match pattern
        match = self.MENTION_PATTERN.match(mention_text)

        if not match:
            return ParsedMention(
                raw_text=mention_text,
                mention_type=MentionType.SIMPLE,
                field_path="",
                parameters={},
                is_valid=False,
                error_message="Invalid mention syntax"
            )

        field_path = match.group(1).lower()  # e.g., "companyname" or "insurance.public_liability"
        params_str = match.group(2)          # e.g., "paul_smith, format=brief" or None

        # Parse parameters if present
        parameters = self._parse_parameters(params_str) if params_str else {}

        # Determine mention type and resolve field path
        mention_type, resolved_path = self._classify_and_resolve(field_path, parameters)

        return ParsedMention(
            raw_text=mention_text,
            mention_type=mention_type,
            field_path=resolved_path,
            parameters=parameters,
            is_valid=True
        )

    def parse_all(self, text: str) -> List[ParsedMention]:
        """
        Find and parse all mentions in text.

        Args:
            text: Document text containing mentions

        Returns:
            List of ParsedMention objects

        Examples:
            >>> parser = MentionParser()
            >>> text = "Company: @companyname, ABN: @abn, Email: @email"
            >>> mentions = parser.parse_all(text)
            >>> len(mentions)
            3
        """
        mentions = []

        for match in self.MENTION_PATTERN.finditer(text):
            mention_text = match.group(0)
            parsed = self.parse(mention_text)
            mentions.append(parsed)

        return mentions

    def _parse_parameters(self, params_str: str) -> Dict[str, str]:
        """
        Parse parameters from bracket content.

        Args:
            params_str: Parameter string (e.g., "paul_smith, format=brief, max_words=500")

        Returns:
            Dict of parameters

        Examples:
            >>> parser = MentionParser()
            >>> parser._parse_parameters("paul_smith")
            {'_positional': 'paul_smith'}

            >>> parser._parse_parameters("paul_smith, format=brief")
            {'_positional': 'paul_smith', 'format': 'brief'}

            >>> parser._parse_parameters("type=executive_summary, length=500")
            {'type': 'executive_summary', 'length': '500'}
        """
        parameters = {}

        # Split by comma
        parts = [p.strip() for p in params_str.split(',')]

        for i, part in enumerate(parts):
            if '=' in part:
                # key=value format
                key, value = part.split('=', 1)
                parameters[key.strip()] = value.strip()
            else:
                # Positional parameter (first param only)
                if i == 0:
                    parameters['_positional'] = part
                else:
                    # Additional positional params go to _args list
                    if '_args' not in parameters:
                        parameters['_args'] = []
                    parameters['_args'].append(part)

        return parameters

    def _classify_and_resolve(
        self,
        field_path: str,
        parameters: Dict[str, str]
    ) -> Tuple[MentionType, str]:
        """
        Classify mention type and resolve field path.

        Args:
            field_path: Raw field path from mention
            parameters: Parsed parameters

        Returns:
            Tuple of (MentionType, resolved_path)

        Examples:
            >>> parser = MentionParser()
            >>> parser._classify_and_resolve("companyname", {})
            (MentionType.SIMPLE, 'company.legal_name')

            >>> parser._classify_and_resolve("cv", {'_positional': 'paul_smith'})
            (MentionType.CONTENT_GEN, 'team.paul_smith')

            >>> parser._classify_and_resolve("generate", {'type': 'executive_summary'})
            (MentionType.CREATIVE_GEN, 'generate.executive_summary')
        """
        root = field_path.split('.')[0]

        # Check if it's a creative generation mention
        if root in self.CREATIVE_TYPES or root == 'generate':
            gen_type = parameters.get('type', root)
            return MentionType.CREATIVE_GEN, f"generate.{gen_type}"

        # Check if it's a simple field (has mapping)
        if root in self.SIMPLE_FIELD_MAPPINGS:
            template = self.SIMPLE_FIELD_MAPPINGS[root]

            # Check if template has placeholders (e.g., team.{person_id})
            if '{' in template:
                # Need to substitute from parameters
                if '_positional' in parameters:
                    # Use positional parameter
                    substituted = template.format(
                        person_id=parameters['_positional'],
                        project_id=parameters['_positional'],
                        reference_id=parameters['_positional'],
                        section=parameters['_positional']
                    )
                    mention_type = MentionType.CONTENT_GEN if root in ['cv', 'project_summary', 'reference'] else MentionType.NARRATIVE
                    return mention_type, substituted
                else:
                    # Missing required parameter
                    return MentionType.SIMPLE, template

            # Direct simple field
            return MentionType.SIMPLE, template

        # Check if it's structured access (has dots)
        if '.' in field_path:
            return MentionType.STRUCTURED, field_path

        # Treat as simple field (use as-is)
        return MentionType.SIMPLE, field_path

    def validate_mention(self, mention: ParsedMention, entity_id: str) -> Tuple[bool, Optional[str]]:
        """
        Validate that a mention can be resolved for given entity.

        Args:
            mention: ParsedMention object
            entity_id: Entity ID to validate against

        Returns:
            Tuple of (is_valid, error_message)

        Note:
            This is a basic validation. Full validation requires EntityProfileManager
            to check if fields actually exist.
        """
        if not mention.is_valid:
            return False, mention.error_message

        # Check required parameters
        if mention.mention_type == MentionType.CONTENT_GEN:
            # Content generation requires ID parameter
            root = mention.root_field

            if root in ['team', 'projects', 'references', 'insurance', 'capabilities']:
                if not mention.parameters.get('_positional'):
                    return False, f"Missing required parameter for {root} mention"

        elif mention.mention_type == MentionType.NARRATIVE:
            # Narrative requires section name
            if not mention.parameters.get('_positional'):
                return False, "Missing section name for narrative mention"

        elif mention.mention_type == MentionType.CREATIVE_GEN:
            # Creative generation requires 'type' parameter if root is 'generate'
            if mention.root_field == 'generate':
                if not mention.parameters.get('type'):
                    return False, "Missing 'type' parameter for generate mention"

        return True, None

    def get_mention_context(self, mention: ParsedMention) -> Dict[str, any]:
        """
        Get context information about a mention (for UI display).

        Args:
            mention: ParsedMention object

        Returns:
            Dict with context information
        """
        context = {
            'type': mention.mention_type.value,
            'field_path': mention.field_path,
            'root': mention.root_field,
            'sub_path': mention.sub_path,
            'parameters': mention.parameters,
            'description': self._get_description(mention)
        }

        return context

    def _get_description(self, mention: ParsedMention) -> str:
        """Generate human-readable description of mention."""
        if mention.mention_type == MentionType.SIMPLE:
            return f"Simple field: {mention.field_path}"

        elif mention.mention_type == MentionType.STRUCTURED:
            return f"Structured data access: {mention.field_path}"

        elif mention.mention_type == MentionType.CONTENT_GEN:
            root = mention.root_field
            item_id = mention.parameters.get('_positional', 'unknown')

            descriptions = {
                'team': f"CV for {item_id}",
                'projects': f"Project summary for {item_id}",
                'references': f"Reference details for {item_id}",
                'insurance': f"Insurance summary for {item_id}",
            }

            return descriptions.get(root, f"Content generation: {mention.field_path}")

        elif mention.mention_type == MentionType.NARRATIVE:
            section = mention.parameters.get('_positional', 'unknown')
            return f"Narrative section: {section}"

        elif mention.mention_type == MentionType.CREATIVE_GEN:
            gen_type = mention.parameters.get('type', mention.root_field)
            return f"LLM-generated content: {gen_type}"

        return "Unknown mention type"


# ============================================
# HELPER FUNCTIONS
# ============================================

def find_mentions_in_text(text: str) -> List[str]:
    """
    Find all mention strings in text (just the strings, not parsed).

    Args:
        text: Document text

    Returns:
        List of mention strings (e.g., ["@companyname", "@abn", ...])

    Examples:
        >>> find_mentions_in_text("Company: @companyname, ABN: @abn")
        ['@companyname', '@abn']
    """
    pattern = re.compile(r'@[a-zA-Z_][a-zA-Z0-9_.]*(?:\[[^\]]+\])?')
    return pattern.findall(text)


def replace_mentions(text: str, replacements: Dict[str, str]) -> str:
    """
    Replace mentions in text with their values.

    Args:
        text: Document text with mentions
        replacements: Dict mapping mention_text → replacement_value

    Returns:
        Text with mentions replaced

    Examples:
        >>> text = "Company: @companyname, ABN: @abn"
        >>> replacements = {"@companyname": "Acme Corp", "@abn": "12 345 678 901"}
        >>> replace_mentions(text, replacements)
        'Company: Acme Corp, ABN: 12 345 678 901'
    """
    result = text

    for mention_text, replacement_value in replacements.items():
        result = result.replace(mention_text, replacement_value)

    return result


def extract_mentions_from_docx(doc_path: str) -> List[Tuple[str, str, int]]:
    """
    Extract all mentions from a .docx file.

    Args:
        doc_path: Path to .docx file

    Returns:
        List of (mention_text, paragraph_text, paragraph_index)

    Examples:
        >>> mentions = extract_mentions_from_docx("tender.docx")
        >>> for mention, para_text, para_idx in mentions:
        ...     print(f"Found {mention} in paragraph {para_idx}")
    """
    import docx

    doc = docx.Document(doc_path)
    parser = MentionParser()
    results = []

    for i, paragraph in enumerate(doc.paragraphs):
        para_text = paragraph.text
        mentions = find_mentions_in_text(para_text)

        for mention in mentions:
            results.append((mention, para_text, i))

    return results
