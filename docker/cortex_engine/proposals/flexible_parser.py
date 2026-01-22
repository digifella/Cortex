"""
Flexible Template Parser - No Rigid Instructions Required
Version: 2.0.0
Date: 2026-01-02

Purpose: Parse tender documents WITHOUT requiring exact [INSTRUCTION] tags.
Detects sections automatically, identifies work needed, enables hint-based assistance.

Key Features:
- Auto-detects sections (headings, numbering, tables)
- Identifies blanks, placeholders, questions
- No format requirements - works with any tender structure
- Hint-based: User can mark any section and ask for help
"""

import re
import docx
from typing import List, Dict, Any, Optional, Tuple
from dataclasses import dataclass, field
from enum import Enum
import logging

logger = logging.getLogger(__name__)


class SectionType(str, Enum):
    """Types of sections detected in tender documents."""
    HEADING = "heading"              # Traditional heading
    NUMBERED = "numbered"            # 1.1, 1.2, etc.
    QUESTION = "question"            # Starts with question words
    REQUIREMENT = "requirement"      # Contains "must", "shall", "required"
    TABLE = "table"                  # Table cell
    BLANK = "blank"                  # Empty or placeholder text
    BOILERPLATE = "boilerplate"      # Already filled, no work needed
    UNKNOWN = "unknown"


class ContentStatus(str, Enum):
    """Status of section content."""
    EMPTY = "empty"                  # No content at all
    PLACEHOLDER = "placeholder"      # Has "[INSERT]", "TBD", "XXX", etc.
    PARTIAL = "partial"              # Some content but seems incomplete
    COMPLETE = "complete"            # Appears to be filled
    UNKNOWN = "unknown"


@dataclass
class FlexibleSection:
    """Represents any section in a tender document - no rigid format required."""

    # Identification
    section_id: str                          # Unique ID (e.g., "section_0", "heading_3")
    section_type: SectionType                # Auto-detected type

    # Content
    heading: str                             # Section heading/title
    content: str                             # Current content (may be empty)
    original_paragraph: Any                  # Reference to docx paragraph

    # Context
    parent_heading: Optional[str] = None     # Parent section (for nested)
    numbering: Optional[str] = None          # e.g., "1.2.3"
    level: int = 0                           # Nesting level (0 = root)

    # Status
    status: ContentStatus = ContentStatus.UNKNOWN
    needs_work: bool = False                 # Auto-detected if section needs filling

    # User hints
    user_hint: Optional[str] = None          # User's guidance for this section
    priority: int = 0                        # User-set priority (0-5)

    # AI assistance metadata
    complexity: str = "moderate"             # simple/moderate/complex
    suggested_approach: Optional[str] = None # What AI recommends

    # Detected patterns
    detected_patterns: List[str] = field(default_factory=list)  # ["question", "requirement", etc.]


    def to_dict(self) -> Dict[str, Any]:
        """Convert to dictionary for storage."""
        return {
            "section_id": self.section_id,
            "section_type": self.section_type.value,
            "heading": self.heading,
            "content": self.content,
            "parent_heading": self.parent_heading,
            "numbering": self.numbering,
            "level": self.level,
            "status": self.status.value,
            "needs_work": self.needs_work,
            "user_hint": self.user_hint,
            "priority": self.priority,
            "complexity": self.complexity,
            "suggested_approach": self.suggested_approach,
            "detected_patterns": self.detected_patterns
        }


class FlexibleTemplateParser:
    """
    Parses tender documents WITHOUT requiring exact [INSTRUCTION] format.
    Works with ANY document structure.
    """

    # Patterns to detect placeholders
    PLACEHOLDER_PATTERNS = [
        r'\[.*?\]',                          # [INSERT], [COMPANY NAME]
        r'<.*?>',                            # <insert text>
        r'_{3,}',                            # _____ (underscores)
        r'\.{3,}',                           # ..... (dots)
        r'\bTBD\b',                          # TBD
        r'\bTBC\b',                          # TBC (To Be Confirmed)
        r'\bXXX\b',                          # XXX
        r'\bTO BE COMPLETED\b',              # TO BE COMPLETED
        r'\bPLEASE COMPLETE\b',              # PLEASE COMPLETE
    ]

    # Question patterns
    QUESTION_PATTERNS = [
        r'^\s*(how|what|when|where|why|who|which|describe|explain|provide|list)',
        r'\?$',                              # Ends with question mark
    ]

    # Requirement patterns
    REQUIREMENT_PATTERNS = [
        r'\b(must|shall|should|required|mandatory)\b',
        r'^\s*\d+\.\s+(must|shall|should)',
    ]

    def __init__(self):
        """Initialize parser."""
        self.sections: List[FlexibleSection] = []
        self.heading_stack: List[Tuple[int, str]] = []  # Track heading hierarchy

    def parse_document(self, doc: docx.document.Document) -> List[FlexibleSection]:
        """
        Parse entire document and detect sections automatically.
        No rigid format required - works with ANY tender structure.
        """
        self.sections = []
        self.heading_stack = []

        paragraphs = list(doc.paragraphs)

        for i, para in enumerate(paragraphs):
            section = self._parse_paragraph(para, i)

            if section:
                # Detect status and needs_work automatically
                self._analyze_section_status(section)

                # Detect complexity
                self._estimate_complexity(section)

                # Suggest approach
                self._suggest_approach(section)

                self.sections.append(section)

        # Parse tables separately
        for i, table in enumerate(doc.tables):
            table_sections = self._parse_table(table, i)
            self.sections.extend(table_sections)

        logger.info(f"Parsed {len(self.sections)} sections from document")
        logger.info(f"Sections needing work: {sum(1 for s in self.sections if s.needs_work)}")

        return self.sections

    def _parse_paragraph(self, para: docx.text.paragraph.Paragraph, index: int) -> Optional[FlexibleSection]:
        """Parse a single paragraph into a FlexibleSection."""

        text = para.text.strip()

        # Skip completely empty paragraphs
        if not text:
            return None

        # Detect section type
        section_type = self._detect_section_type(para, text)

        # Extract heading and content
        heading, content = self._extract_heading_and_content(para, text, section_type)

        # Detect numbering (e.g., "1.2.3")
        numbering = self._extract_numbering(text)

        # Determine nesting level
        level = self._get_nesting_level(para, section_type)

        # Update heading stack
        if section_type == SectionType.HEADING:
            self._update_heading_stack(level, heading)

        # Get parent heading
        parent = self._get_parent_heading(level)

        # Create section
        section = FlexibleSection(
            section_id=f"section_{index}",
            section_type=section_type,
            heading=heading,
            content=content,
            original_paragraph=para,
            parent_heading=parent,
            numbering=numbering,
            level=level
        )

        return section

    def _detect_section_type(self, para: docx.text.paragraph.Paragraph, text: str) -> SectionType:
        """Auto-detect section type from paragraph."""

        # Check if it's a heading style
        if para.style.name.startswith('Heading'):
            return SectionType.HEADING

        # Check for numbering (1.1, 1.2, etc.)
        if re.match(r'^\s*\d+(\.\d+)*\.?\s+', text):
            return SectionType.NUMBERED

        # Check for questions
        if any(re.search(pattern, text, re.IGNORECASE) for pattern in self.QUESTION_PATTERNS):
            return SectionType.QUESTION

        # Check for requirements
        if any(re.search(pattern, text, re.IGNORECASE) for pattern in self.REQUIREMENT_PATTERNS):
            return SectionType.REQUIREMENT

        # Check if blank/placeholder
        if self._is_placeholder(text):
            return SectionType.BLANK

        # Default to boilerplate if it has substantive content
        if len(text) > 50:
            return SectionType.BOILERPLATE

        return SectionType.UNKNOWN

    def _extract_heading_and_content(
        self,
        para: docx.text.paragraph.Paragraph,
        text: str,
        section_type: SectionType
    ) -> Tuple[str, str]:
        """Extract heading and content from paragraph text."""

        if section_type == SectionType.HEADING:
            # Entire text is the heading
            return text, ""

        if section_type == SectionType.NUMBERED:
            # Split on numbering
            match = re.match(r'^\s*(\d+(?:\.\d+)*\.?)\s+(.*)', text)
            if match:
                numbering, rest = match.groups()
                # First sentence might be heading
                sentences = rest.split('.')
                if sentences:
                    heading = sentences[0].strip()
                    content = '.'.join(sentences[1:]).strip() if len(sentences) > 1 else ""
                    return heading, content

        if section_type == SectionType.QUESTION:
            # Question is the heading
            return text, ""

        # For other types, no clear heading
        return "Content", text

    def _extract_numbering(self, text: str) -> Optional[str]:
        """Extract numbering like '1.2.3' from text."""
        match = re.match(r'^\s*(\d+(?:\.\d+)*)', text)
        return match.group(1) if match else None

    def _get_nesting_level(self, para: docx.text.paragraph.Paragraph, section_type: SectionType) -> int:
        """Determine nesting level from paragraph."""

        # Heading styles have levels
        if para.style.name.startswith('Heading'):
            try:
                return int(para.style.name.replace('Heading', '').strip() or '1') - 1
            except ValueError:
                return 0

        # Numbering depth (1 = level 0, 1.1 = level 1, 1.1.1 = level 2)
        if section_type == SectionType.NUMBERED:
            numbering = self._extract_numbering(para.text)
            if numbering:
                return numbering.count('.')

        return 0

    def _update_heading_stack(self, level: int, heading: str):
        """Update heading hierarchy stack."""
        # Remove any headings at same or deeper level
        self.heading_stack = [(l, h) for l, h in self.heading_stack if l < level]
        # Add current heading
        self.heading_stack.append((level, heading))

    def _get_parent_heading(self, level: int) -> Optional[str]:
        """Get parent heading for current level."""
        # Find most recent heading at shallower level
        for l, h in reversed(self.heading_stack):
            if l < level:
                return h
        return None

    def _is_placeholder(self, text: str) -> bool:
        """Check if text contains placeholders."""
        return any(re.search(pattern, text, re.IGNORECASE) for pattern in self.PLACEHOLDER_PATTERNS)

    def _analyze_section_status(self, section: FlexibleSection):
        """Automatically determine if section needs work."""

        content = section.content.strip()

        # Empty content
        if not content:
            section.status = ContentStatus.EMPTY
            section.needs_work = True
            section.detected_patterns.append("empty")
            return

        # Has placeholder text
        if self._is_placeholder(content):
            section.status = ContentStatus.PLACEHOLDER
            section.needs_work = True
            section.detected_patterns.append("placeholder")
            return

        # Very short content (likely incomplete)
        if len(content) < 50:
            section.status = ContentStatus.PARTIAL
            section.needs_work = True
            section.detected_patterns.append("too_short")
            return

        # Questions without answers
        if section.section_type == SectionType.QUESTION and len(content) < 100:
            section.status = ContentStatus.PARTIAL
            section.needs_work = True
            section.detected_patterns.append("unanswered_question")
            return

        # Appears complete
        section.status = ContentStatus.COMPLETE
        section.needs_work = False

    def _estimate_complexity(self, section: FlexibleSection):
        """Estimate section complexity for MoE routing."""

        # Simple: Boilerplate, already complete, short
        if section.status == ContentStatus.COMPLETE and len(section.heading) < 50:
            section.complexity = "simple"
            return

        # Complex: Questions, requirements, technical sections
        complex_keywords = [
            'technical', 'methodology', 'approach', 'strategy', 'innovation',
            'describe how', 'explain your', 'detail your'
        ]

        combined_text = f"{section.heading} {section.content}".lower()

        if any(keyword in combined_text for keyword in complex_keywords):
            section.complexity = "complex"
            return

        # Default to moderate
        section.complexity = "moderate"

    def _suggest_approach(self, section: FlexibleSection):
        """Suggest AI approach based on section characteristics."""

        if not section.needs_work:
            section.suggested_approach = "No action needed - section is complete"
            return

        if section.section_type == SectionType.QUESTION:
            section.suggested_approach = "Answer the question using knowledge base and expertise"
            return

        if section.section_type == SectionType.REQUIREMENT:
            section.suggested_approach = "Address the requirement with specific capabilities"
            return

        if "methodology" in section.heading.lower():
            section.suggested_approach = "Generate detailed technical approach with MoE"
            return

        if section.status == ContentStatus.EMPTY:
            section.suggested_approach = "Generate content from knowledge base"
            return

        if section.status == ContentStatus.PARTIAL:
            section.suggested_approach = "Refine and expand existing content"
            return

        section.suggested_approach = "General content generation"

    def _parse_table(self, table: docx.table.Table, table_index: int) -> List[FlexibleSection]:
        """Parse table into sections (each cell can be a section)."""

        sections = []

        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                text = cell.text.strip()

                if not text:
                    continue

                # Check if this cell needs work
                needs_work = self._is_placeholder(text) or len(text) < 20

                section = FlexibleSection(
                    section_id=f"table_{table_index}_r{row_idx}_c{col_idx}",
                    section_type=SectionType.TABLE,
                    heading=f"Table {table_index + 1}, Row {row_idx + 1}, Col {col_idx + 1}",
                    content=text,
                    original_paragraph=cell,
                    level=0,
                    needs_work=needs_work,
                    status=ContentStatus.PLACEHOLDER if needs_work else ContentStatus.COMPLETE
                )

                sections.append(section)

        return sections

    def get_sections_needing_work(self) -> List[FlexibleSection]:
        """Get all sections that need work."""
        return [s for s in self.sections if s.needs_work]

    def get_sections_by_priority(self) -> List[FlexibleSection]:
        """Get sections sorted by user priority (highest first)."""
        return sorted(self.sections, key=lambda s: s.priority, reverse=True)

    def get_section_by_id(self, section_id: str) -> Optional[FlexibleSection]:
        """Get specific section by ID."""
        for section in self.sections:
            if section.section_id == section_id:
                return section
        return None
